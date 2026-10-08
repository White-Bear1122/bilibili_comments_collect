# -*- coding: utf-8 -*-
"""
Bilibili 评论爬取网页端

功能：
  - 输入视频 BV 号，分页列出评论及楼中楼回复
  - 按用户名筛选
  - 收集单条评论及其全部回复为 JSON
  - 导出当前已加载评论为 Excel

运行：
  pip install -r requirements.txt
  python bilibili_comments_web.py

访问：http://127.0.0.1:8000
"""

import json
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path

import requests
from fastapi import FastAPI, Form
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "templates" / "index.html"
STATIC_DIR = BASE_DIR / "static"
SAVE_DIR = BASE_DIR / "collected"
SAVE_DIR.mkdir(exist_ok=True)

app = FastAPI(title="Bilibili 评论爬取工具")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
    ),
    "Referer": "https://www.bilibili.com",
    "Origin": "https://www.bilibili.com",
    "Accept": "application/json, text/plain, */*",
}

# 当前会话已加载的评论（服务端内存，单用户本地使用足够）
state = {
    "bv": "",
    "title": "",
    "oid": 0,
    "page": 0,
    "page_size": 20,
    "has_more": False,
    "comments": [],
}


def get_video_info(bvid: str):
    try:
        url = f"https://api.bilibili.com/x/web-interface/view?bvid={bvid}"
        r = requests.get(url, headers=HEADERS, timeout=10)
        data = r.json()
        if data.get("code") == 0:
            return {
                "code": 0,
                "title": data["data"]["title"],
                "aid": data["data"]["aid"],
            }
        return {"code": data.get("code"), "message": data.get("message", "获取视频信息失败")}
    except Exception as e:
        return {"code": -1, "message": str(e)}


def parse_replies(replies_list):
    subs = []
    for sub in replies_list or []:
        member = sub.get("member") or {}
        subs.append(
            {
                "rpid": sub.get("rpid"),
                "user": member.get("uname") or member.get("name", ""),
                "content": (sub.get("content") or {}).get("message", ""),
                "like": sub.get("like", 0),
                "ctime": datetime.fromtimestamp(sub.get("ctime", 0)).strftime("%Y-%m-%d %H:%M"),
            }
        )
    return subs


def fetch_all_sub_replies(oid: int, root_rpid: int):
    """拉取某条主评论下的全部楼中楼。"""
    collected = []
    page = 1
    while page <= 50:
        try:
            url = "https://api.bilibili.com/x/v2/reply/reply"
            params = {"type": 1, "oid": oid, "root": root_rpid, "pn": page, "ps": 20}
            r = requests.get(url, params=params, headers=HEADERS, timeout=12)
            data = r.json()
            if data.get("code") != 0:
                break
            replies = (data.get("data") or {}).get("replies") or []
            if not replies:
                break
            collected.extend(parse_replies(replies))
            if len(replies) < 20:
                break
            page += 1
            time.sleep(0.3)
        except Exception:
            break
    return collected


def fetch_page(oid: int, page: int, size: int = 20):
    """拉取一页主评论。楼中楼优先用接口已返回的 replies，不足时再补拉。"""
    try:
        url = "https://api.bilibili.com/x/v2/reply/main"
        params = {"type": 1, "oid": oid, "mode": 3, "next": page - 1, "ps": size}
        r = requests.get(url, params=params, headers=HEADERS, timeout=12)
        data = r.json()
        if data.get("code") != 0:
            url2 = "https://api.bilibili.com/x/v2/reply"
            params2 = {"type": 1, "oid": oid, "sort": 2, "pn": page, "ps": size}
            r = requests.get(url2, params=params2, headers=HEADERS, timeout=12)
            data = r.json()
            if data.get("code") != 0:
                return {"code": data.get("code"), "message": data.get("message"), "comments": [], "has_more": False}
            replies = (data.get("data") or {}).get("replies") or []
            cursor_more = len(replies) == size
        else:
            body = data.get("data") or {}
            replies = body.get("replies") or []
            cursor = body.get("cursor") or {}
            cursor_more = not cursor.get("is_end", len(replies) < size)

        comments = []
        for reply in replies:
            member = reply.get("member") or {}
            comment = {
                "rpid": reply.get("rpid"),
                "user": member.get("uname") or member.get("name", ""),
                "uid": member.get("mid"),
                "content": (reply.get("content") or {}).get("message", ""),
                "like": reply.get("like", 0),
                "ctime": datetime.fromtimestamp(reply.get("ctime", 0)).strftime("%Y-%m-%d %H:%M"),
                "reply_count": reply.get("rcount", 0),
                "replies": parse_replies(reply.get("replies")),
            }
            if comment["reply_count"] > len(comment["replies"]):
                comment["replies"] = fetch_all_sub_replies(oid, comment["rpid"]) or comment["replies"]
            comments.append(comment)

        return {"code": 0, "comments": comments, "has_more": bool(cursor_more and replies)}
    except Exception as e:
        return {"code": -1, "message": str(e), "comments": [], "has_more": False}


def filter_by_user(comments, username: str):
    name = (username or "").strip().lower()
    if not name:
        return comments
    result = []
    for comment in comments:
        hit_main = name in (comment.get("user") or "").lower()
        hit_sub = any(name in (sub.get("user") or "").lower() for sub in comment.get("replies") or [])
        if hit_main or hit_sub:
            result.append(comment)
    return result


def comments_to_workbook(comments, title: str, bvid: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "评论"

    headers = ["层级", "主评论ID", "用户", "UID", "点赞", "时间", "内容"]
    header_fill = PatternFill("solid", fgColor="FB7299")
    header_font = Font(bold=True, color="FFFFFF")
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row = 2
    for comment in comments:
        ws.cell(row, 1, "主评论")
        ws.cell(row, 2, comment.get("rpid"))
        ws.cell(row, 3, comment.get("user"))
        ws.cell(row, 4, comment.get("uid") or "")
        ws.cell(row, 5, comment.get("like", 0))
        ws.cell(row, 6, comment.get("ctime"))
        content_cell = ws.cell(row, 7, comment.get("content"))
        content_cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
        for sub in comment.get("replies") or []:
            ws.cell(row, 1, "回复")
            ws.cell(row, 2, comment.get("rpid"))
            ws.cell(row, 3, sub.get("user"))
            ws.cell(row, 4, "")
            ws.cell(row, 5, sub.get("like", 0))
            ws.cell(row, 6, sub.get("ctime"))
            sub_cell = ws.cell(row, 7, sub.get("content"))
            sub_cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

    widths = [12, 18, 18, 16, 10, 20, 80]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.auto_filter.ref = f"A1:G{max(row - 1, 1)}"
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "FB7299"

    meta = wb.create_sheet("说明", 0)
    meta["A1"] = "视频标题"
    meta["B1"] = title
    meta["A2"] = "BV号"
    meta["B2"] = bvid
    meta["A3"] = "导出时间"
    meta["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta["A4"] = "主评论条数"
    meta["B4"] = len(comments)
    meta["A5"] = "回复条数"
    meta["B5"] = sum(len(comment.get("replies") or []) for comment in comments)
    meta.column_dimensions["A"].width = 16
    meta.column_dimensions["B"].width = 60
    return wb


def normalize_bvid(raw: str) -> str:
    value = (raw or "").strip()
    if value.lower().startswith("http"):
        path = value.split("?")[0].rstrip("/")
        value = path.split("/")[-1]
    return value


@app.get("/", response_class=HTMLResponse)
async def home():
    return HTMLResponse(TEMPLATE_PATH.read_text(encoding="utf-8"))


@app.post("/api/fetch")
async def api_fetch(bv: str = Form(...), reset: str = Form("1")):
    bv = normalize_bvid(bv)
    if not bv.upper().startswith("BV"):
        return JSONResponse({"status": "error", "message": "请输入有效的 BV 号或视频链接"})

    if reset == "1" or state["bv"] != bv:
        info = get_video_info(bv)
        if info.get("code") != 0:
            return JSONResponse({"status": "error", "message": info.get("message", "视频不存在")})
        state.update(
            {
                "bv": bv,
                "title": info["title"],
                "oid": info["aid"],
                "page": 0,
                "has_more": True,
                "comments": [],
            }
        )

    if state["page"] == 0:
        page_data = fetch_page(state["oid"], 1, state["page_size"])
        if page_data.get("code") != 0:
            return JSONResponse({"status": "error", "message": page_data.get("message", "评论获取失败")})
        state["comments"] = page_data["comments"]
        state["page"] = 1
        state["has_more"] = page_data["has_more"]

    return JSONResponse(
        {
            "status": "ok",
            "title": state["title"],
            "bv": state["bv"],
            "comments": state["comments"],
            "total": len(state["comments"]),
            "has_more": state["has_more"],
        }
    )


@app.post("/api/more")
async def api_more():
    if not state["oid"]:
        return JSONResponse({"status": "error", "message": "请先输入 BV 号"})
    if not state["has_more"]:
        return JSONResponse({"status": "error", "message": "没有更多评论了"})

    next_page = state["page"] + 1
    page_data = fetch_page(state["oid"], next_page, state["page_size"])
    if page_data.get("code") != 0:
        return JSONResponse({"status": "error", "message": page_data.get("message", "加载失败")})

    seen = {comment["rpid"] for comment in state["comments"]}
    for comment in page_data["comments"]:
        if comment["rpid"] not in seen:
            state["comments"].append(comment)
    state["page"] = next_page
    state["has_more"] = page_data["has_more"] and bool(page_data["comments"])

    return JSONResponse(
        {
            "status": "ok",
            "title": state["title"],
            "bv": state["bv"],
            "comments": state["comments"],
            "total": len(state["comments"]),
            "has_more": state["has_more"],
        }
    )


@app.post("/api/filter")
async def api_filter(username: str = Form("")):
    if not state["comments"]:
        return JSONResponse({"status": "error", "message": "还没有加载评论"})
    return JSONResponse({"status": "ok", "comments": filter_by_user(state["comments"], username)})


@app.post("/api/collect")
async def api_collect(rpid: str = Form(...)):
    target = next((comment for comment in state["comments"] if str(comment.get("rpid")) == str(rpid)), None)
    if not target:
        return JSONResponse({"status": "error", "message": "未找到该评论，请先加载到这一页"})

    if state["oid"] and target.get("rpid"):
        full = fetch_all_sub_replies(state["oid"], target["rpid"])
        if full:
            target["replies"] = full

    payload = {
        "bv": state["bv"],
        "title": state["title"],
        "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "comment": target,
    }
    filename = f"collected_{state['bv']}_{target['rpid']}.json"
    path = SAVE_DIR / filename
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return JSONResponse(
        {
            "status": "ok",
            "message": f"已保存该评论及 {len(target.get('replies') or [])} 条回复：collected/{filename}",
        }
    )


@app.post("/api/export")
async def api_export(username: str = Form("")):
    if not state["comments"]:
        return JSONResponse({"status": "error", "message": "还没有可导出的评论"}, status_code=400)

    rows = filter_by_user(state["comments"], username) if username.strip() else state["comments"]
    if not rows:
        return JSONResponse({"status": "error", "message": "筛选结果为空，无法导出"}, status_code=400)

    workbook = comments_to_workbook(rows, state["title"], state["bv"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"bilibili_{state['bv']}_{stamp}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename={filename}"}
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.get("/collected/{name}")
async def download_collected(name: str):
    path = (SAVE_DIR / name).resolve()
    if SAVE_DIR.resolve() not in path.parents or not path.is_file():
        return JSONResponse({"status": "error", "message": "文件不存在"}, status_code=404)
    return FileResponse(path, filename=path.name)


if __name__ == "__main__":
    import uvicorn

    print("Bilibili 评论爬取网页已启动: http://127.0.0.1:8000")
    uvicorn.run(app, host="127.0.0.1", port=8000)
